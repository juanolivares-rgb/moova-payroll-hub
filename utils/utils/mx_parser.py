from openpyxl import load_workbook
import re

def parse_mx_ops(file_obj, manuales=None):
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = list(wb.active.iter_rows(values_only=True))
    tope = 9
    periodo = ''

    r0 = ws[0] if ws else []
    for c in r0:
        s = str(c or '')
        if 'Tope' in s:
            m = re.search(r'(\d+)', s)
            if m: tope = int(m.group(1))
        if '/' in s and len(s) > 8: periodo = s.strip()

    hdr = next((i for i, r in enumerate(ws) if r and r[1] == 'Colaborador'), -1)
    colabs = {}

    if hdr >= 0:
        for r in ws[hdr + 1:]:
            nombre = str(r[1] or '').strip()
            if not nombre: continue
            tipo = str(r[2] or '').strip()
            fecha = str(r[3] or '')
            dia = str(r[5] or '')
            horario = str(r[6] or '')
            diasDesc = float(r[8] or 0)
            heDoble = float(r[9] or 0)
            heTriple = float(r[10] or 0)
            heDom = float(r[11] or 0)
            heFest = float(r[12] or 0)
            primaDom = float(r[13] or 0)
            comentarios = str(r[14] or '') if len(r) > 14 else ''

            if nombre not in colabs:
                colabs[nombre] = {
                    'heDoble': 0, 'heTriple': 0, 'heDom': 0, 'heFest': 0,
                    'diasDesc': 0, 'primaDom': 0, 'comentarios': '',
                    'alertas': [], 'filas': []
                }
            colabs[nombre]['heDoble'] += heDoble
            colabs[nombre]['heTriple'] += heTriple
            colabs[nombre]['heDom'] += heDom
            colabs[nombre]['heFest'] += heFest
            colabs[nombre]['diasDesc'] += diasDesc
            colabs[nombre]['primaDom'] += primaDom
            if comentarios and not colabs[nombre]['comentarios']:
                colabs[nombre]['comentarios'] = comentarios
            colabs[nombre]['filas'].append({
                'tipo': tipo, 'fecha': fecha, 'dia': dia, 'horario': horario,
                'heDoble': heDoble, 'heTriple': heTriple,
                'heDom': heDom, 'heFest': heFest,
                'diasDesc': diasDesc, 'primaDom': primaDom,
                'comentarios': comentarios
            })

    # Novedades manuales
    if manuales:
        for m in manuales:
            nombre = m.get('nombre', '').strip()
            if not nombre: continue
            if nombre not in colabs:
                colabs[nombre] = {
                    'heDoble': 0, 'heTriple': 0, 'heDom': 0, 'heFest': 0,
                    'diasDesc': 0, 'primaDom': 0, 'comentarios': '',
                    'alertas': [], 'filas': []
                }
            tipo = m.get('tipo', '')
            val = float(m.get('valor', 0))
            if tipo == 'heDoble': colabs[nombre]['heDoble'] += val
            elif tipo == 'heTriple': colabs[nombre]['heTriple'] += val
            elif tipo == 'heDom': colabs[nombre]['heDom'] += val
            elif tipo == 'diasDesc': colabs[nombre]['diasDesc'] += val
            elif tipo == 'bono': colabs[nombre].setdefault('bono', 0); colabs[nombre]['bono'] += val
            elif tipo == 'ausencia': colabs[nombre].setdefault('ausencia', 0); colabs[nombre]['ausencia'] += val

    # Validaciones LFT
    for nombre, d in colabs.items():
        totalHE = d['heDoble'] + d['heTriple'] + d['heDom'] + d['heFest']
        esIncap = 'INCAPACIDAD' in d['comentarios'].upper() or any(
            'INCAPACIDAD' in f['comentarios'].upper() for f in d['filas'])
        if esIncap and totalHE > 0:
            d['alertas'].append({
                'tipo': 'err',
                'msg': 'HH.EE registradas durante incapacidad — no liquidable (LFT Art. 42)'
            })
        if totalHE > tope * 2:
            d['alertas'].append({
                'tipo': 'warn',
                'msg': f'Total HH.EE ({totalHE}h) puede superar tope semanal de {tope}h — verificar distribución'
            })
        d['totalHE'] = totalHE

    return colabs, periodo, tope


def parse_mx_prenomina(file_obj):
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws_name = next((s for s in wb.sheetnames if 'NOMINA' in s.upper()), wb.sheetnames[0])
    rows = list(wb[ws_name].iter_rows(values_only=True))
    empleados = {}
    emp_actual = None
    periodo = ''

    for r in rows:
        c0 = str(r[0] or '').strip()
        if 'Período del' in c0: periodo = c0.strip()
        if c0 and c0.isdigit() and len(c0) == 6:
            nombre = ' '.join(str(r[1] or '').split())
            emp_actual = nombre
            empleados[nombre] = {
                'clave': c0,
                'totalPerc': float(r[5] or 0),
                'totalISR': float(r[7] or 0),
                'totalDesc': float(r[8] or 0),
                'neto': float(r[9] or 0),
                'conceptos': []
            }
        elif emp_actual and r[5] and str(r[5]).strip().isdigit():
            empleados[emp_actual]['conceptos'].append({
                'cod': str(r[5]).strip(),
                'desc': str(r[6] or '').strip(),
                'cantidad': r[7],
                'perc': float(r[8] or 0) if r[8] else None,
                'ded': float(r[9] or 0) if r[9] else None
            })

    return empleados, periodo


def run_mx_control(colabs, empleados):
    def norm(s): return ' '.join(s.upper().split())

    results = []
    for nombreEst, emp in empleados.items():
        opsKey = next((k for k in colabs if any(
            w in norm(nombreEst) for w in norm(k).split() if len(w) > 3)), None)
        ops = colabs.get(opsKey)

        def getCant(cod):
            return sum(float(c['cantidad'] or 0) for c in emp['conceptos']
                       if c['cod'] == cod and c['cantidad'])

        heDobleEst = getCant('028')
        heTripleEst = getCant('029')
        diasDescEst = getCant('030')
        primaDomEst = getCant('033')

        heDobleOPS = ops['heDoble'] if ops else None
        heTripleOPS = ops['heTriple'] if ops else None
        diasDescOPS = ops['diasDesc'] if ops else None
        primaDomOPS = ops['primaDom'] if ops else None

        dHD = round(heDobleEst - (heDobleOPS or 0), 2) if heDobleOPS is not None else None
        dHT = round(heTripleEst - (heTripleOPS or 0), 2) if heTripleOPS is not None else None
        dDD = round(diasDescEst - (diasDescOPS or 0), 2) if diasDescOPS is not None else None
        dPD = round(primaDomEst - (primaDomOPS or 0), 2) if primaDomOPS is not None else None

        hasDiff = any(d is not None and abs(d) > 0.1 for d in [dHD, dHT, dDD, dPD])
        esIncap = ops and any('incapacidad' in a['msg'].lower() for a in ops.get('alertas', []))
        tieneHEliq = heDobleEst > 0 or heTripleEst > 0
        alertaIncap = esIncap and tieneHEliq

        results.append({
            'nombre': nombreEst,
            'clave': emp['clave'],
            'opsMatch': opsKey,
            'heDobleOPS': heDobleOPS,
            'heTripleOPS': heTripleOPS,
            'diasDescOPS': diasDescOPS,
            'primaDomOPS': primaDomOPS,
            'heDobleEst': heDobleEst,
            'heTripleEst': heTripleEst,
            'diasDescEst': diasDescEst,
            'primaDomEst': primaDomEst,
            'dHD': dHD, 'dHT': dHT, 'dDD': dDD, 'dPD': dPD,
            'totalPerc': emp['totalPerc'],
            'totalDesc': emp['totalDesc'],
            'neto': emp['neto'],
            'hasDiff': hasDiff,
            'alertaIncap': alertaIncap,
        })

    return sorted(results, key=lambda x: (-x['hasDiff'], x['nombre']))
